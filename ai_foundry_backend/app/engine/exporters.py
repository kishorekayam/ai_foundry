"""File builders: Excel, CSV ZIP, JSON, SQL scripts, Postman v2.1, Pytest, test-case workbook, reports."""
from __future__ import annotations

import csv
import io
import json
import re
import zipfile
from datetime import datetime, timezone
from typing import Any, Callable

from openpyxl import Workbook

from ..models import Report, Schema, TestCase
from .generator import Dataset

STAMP = "atom_foundry"
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

FORMATS = {"xlsx", "csv-zip", "json", "sql-postgres", "sql-mysql", "postman", "testcase-xlsx", "pytest", "mapping-report", "report-json", "report-xlsx"}


def _cell(v: Any) -> Any:
    return json.dumps(v, default=str) if isinstance(v, (list, dict)) else v


def _sheet(wb: Workbook, title: str, rows: list[dict[str, Any]], first: bool = False) -> None:
    ws = wb.active if first else wb.create_sheet()
    ws.title = re.sub(r"[\[\]\*\?/\\:]", "_", title)[:31]
    cols = list(rows[0].keys()) if rows else []
    ws.append(cols)
    for r in rows:
        ws.append([("" if r.get(c) is None else _cell(r.get(c))) for c in cols])


def _xlsx_bytes(sheets: list[tuple[str, list[dict[str, Any]]]]) -> bytes:
    wb = Workbook()
    for i, (title, rows) in enumerate(sheets):
        _sheet(wb, title, rows, first=i == 0)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _csv(rows: list[dict[str, Any]], cols: list[str]) -> str:
    buf = io.StringIO()
    w = csv.writer(buf, lineterminator="\n")
    w.writerow(cols)
    for r in rows:
        w.writerow(["" if r.get(c) is None else _cell(r.get(c)) for c in cols])
    return buf.getvalue()


def _lit(v: Any) -> str:
    if v is None:
        return "NULL"
    if isinstance(v, bool):
        return "TRUE" if v else "FALSE"
    if isinstance(v, (int, float)):
        return str(v)
    if isinstance(v, list):
        return "'{" + ",".join('"' + str(x).replace('"', '\\"') + '"' for x in v) + "}'"
    return "'" + str(v).replace("'", "''") + "'"


def _sql(schema: Schema, dataset: Dataset, order: list[str], q: Callable[[str], str], dialect: str, begin: str) -> str:
    sch = schema.schema_name if dialect == "PostgreSQL" and schema.schema_name and schema.schema_name != "public" else None
    qual = (lambda t: f"{q(sch)}.{q(t)}") if sch else q
    lines = [f"-- ATOM DATA FOUNDRY synthetic data ({dialect})", f"-- Database: {schema.database}" + (f", schema: {sch}" if sch else ""), f"-- Generated: {datetime.now(timezone.utc).isoformat()}", begin, ""]
    for t in order:
        rows = dataset.get(t) or []
        if not rows:
            continue
        cols = list(rows[0].keys())
        lines.append(f"-- {t}: {len(rows)} rows")
        for k in range(0, len(rows), 100):
            chunk = rows[k : k + 100]
            lines.append(f"INSERT INTO {qual(t)} ({', '.join(q(c) for c in cols)}) VALUES")
            lines.append(",\n".join("  (" + ", ".join(_lit(r.get(c)) for c in cols) + ")" for r in chunk) + ";")
        lines.append("")
    lines.append("COMMIT;")
    return "\n".join(lines)


def _body(fx: Any) -> str:
    return json.dumps(fx.body, indent=2, default=str) if fx is not None and fx.body is not None else ""


def _case_rows(cases: list[TestCase]) -> list[dict[str, Any]]:
    return [
        {
            "Test ID": t.id, "Category": t.category, "Title": t.title, "Endpoint": t.endpoint or "", "Table": t.table,
            "Preconditions": "\n".join(t.preconditions), "Steps": "\n".join(f"{i + 1}. {s}" for i, s in enumerate(t.steps)),
            "Expected result": t.expected, "Expected status": (t.fixture.expected_status if t.fixture and not t.requests else " -> ".join(str(r.expected_status) for r in t.requests or [])),
            "Request body": _body(t.fixture), "Needs rules": "Yes" if t.needs_rules else "No", "Ready": "Ready" if t.ready else "Needs input",
        }
        for t in cases
    ]


def _coverage_rows(cases: list[TestCase]) -> list[dict[str, Any]]:
    rows: dict[str, dict[str, Any]] = {}
    for t in cases:
        r = rows.setdefault(t.endpoint or "(none)", {"Endpoint": t.endpoint or "(none)", "Table": t.table or "(not matched)", "positive": 0, "negative": 0, "boundary": 0, "workflow": 0})
        r[t.category] += 1
    return [{**r, "Total": r["positive"] + r["negative"] + r["boundary"] + r["workflow"]} for r in rows.values()]


def _status_script(expected: str | None) -> str:
    e = expected or "2xx"
    if e.endswith("xx"):
        lo = int(e[0]) * 100
        return f'pm.test("Status is {e}", () => pm.expect(pm.response.code).to.be.within({lo}, {lo + 99}));'
    return f'pm.test("Status is {e}", () => pm.expect(pm.response.code).to.eql({int(e)}));'


def _pm_request(name: str, fx: Any) -> dict[str, Any]:
    path = re.sub(r"(?<!\{)\{([A-Za-z_][\w-]*)\}(?!\})", r"{{\g<1>}}", fx.path)  # leftover {id} -> Postman variable
    headers = [{"key": k, "value": v} for k, v in fx.headers.items()]
    if fx.body is not None and not any(h["key"].lower() == "content-type" for h in headers):
        headers.append({"key": "Content-Type", "value": "application/json"})
    exec_: list[str] = [_status_script(fx.expected_status)]
    if fx.note and fx.note.startswith("capture-id:"):
        pk = fx.note.split(":", 1)[1]
        exec_.append(f'try {{ const j = pm.response.json(); const d = j.data || j; pm.collectionVariables.set("createdId", d["{pk}"] ?? d.id ?? d.pk); }} catch (e) {{}}')
    req: dict[str, Any] = {"method": fx.method, "header": headers, "url": {"raw": "{{baseUrl}}" + path, "host": ["{{baseUrl}}"], "path": [p for p in path.split("/") if p]}}
    if fx.body is not None:
        req["body"] = {"mode": "raw", "raw": json.dumps(fx.body, indent=2, default=str), "options": {"raw": {"language": "json"}}}
    return {"name": name, "request": req, "event": [{"listen": "test", "script": {"type": "text/javascript", "exec": exec_}}]}


def _postman(schema: Schema, cases: list[TestCase], api_base: str | None) -> dict[str, Any]:
    folders: dict[str, list[dict[str, Any]]] = {}
    for t in cases:
        if not t.fixture:
            continue
        title = f"{t.id} {t.title}"
        if t.requests:
            item: dict[str, Any] = {"name": title, "item": [_pm_request(f"Step {i + 1}: {r.method} {r.path}", r) for i, r in enumerate(t.requests)]}
        else:
            item = _pm_request(title, t.fixture)
        folders.setdefault(t.category.capitalize(), []).append(item)
    return {
        "info": {"name": f"ATOM DATA FOUNDRY - test cases ({schema.database})", "schema": "https://schema.getpostman.com/json/collection/v2.1.0/collection.json"},
        "variable": [{"key": "baseUrl", "value": api_base or "http://localhost:8000"}, {"key": "authToken", "value": ""}, {"key": "createdId", "value": ""}],
        "item": [{"name": k, "item": v} for k, v in folders.items()],
    }


_PYTEST_HEADER = r'''import os
import re

import httpx

BASE_URL = os.environ.get("BASE_URL", "http://localhost:8000")
AUTH_TOKEN = os.environ.get("AUTH_TOKEN", "")


def _headers(h):
    return {k: v.replace("{{authToken}}", AUTH_TOKEN) for k, v in h.items()}


def _expect(r, status):
    if status.endswith("xx"):
        lo = int(status[0]) * 100
        assert lo <= r.status_code < lo + 100, f"{r.status_code}: {r.text[:200]}"
    else:
        assert r.status_code == int(status), f"{r.status_code}: {r.text[:200]}"


def _send(method, path, body, headers, status, created=None):
    path = path.replace("{{createdId}}", str(created))
    path = re.sub(r"\{\{?([A-Za-z_][\w-]*)\}?\}", lambda m: os.environ.get(m.group(1).upper(), "1"), path)  # other path variables come from the environment
    r = httpx.request(method, BASE_URL + path, json=body, headers=_headers(headers), timeout=30)
    _expect(r, status)
    return r

'''


def _pytest(cases: list[TestCase]) -> str:
    out = [_PYTEST_HEADER]
    for t in cases:
        if not t.fixture:
            continue
        slug = re.sub(r"[^a-z0-9]+", "_", t.title.lower()).strip("_")[:48]
        out += [f"def test_{t.id.lower()}_{slug}():", f'    """{t.title.replace(chr(34), chr(39))}. Expected: {t.expected.replace(chr(34), chr(39))}"""', "    created = None"]
        for r in t.requests or [t.fixture]:
            call = f"_send({r.method!r}, {r.path!r}, {r.body!r}, {r.headers!r}, {r.expected_status or '2xx'!r}, created)"
            if r.note and r.note.startswith("capture-id:"):
                pk = r.note.split(":", 1)[1]
                out += [f"    resp = {call}", "    j = resp.json()", f'    created = (j.get("data") or j).get("{pk}", (j.get("data") or j).get("id"))']
            else:
                out.append(f"    {call}")
        out.append("")
    return "\n".join(out)


def build_export(fmt: str, schema: Schema, dataset: Dataset, order: list[str], report: Report, cases: list[TestCase] | None, api_base: str | None) -> tuple[bytes, str, str]:
    """Returns (content, filename, media_type)."""
    names = [t for t in order if dataset.get(t) is not None]
    cases = cases or []
    if fmt == "xlsx":
        return _xlsx_bytes([(t, dataset[t]) for t in names]), f"{STAMP}_data.xlsx", XLSX
    if fmt == "csv-zip":
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
            for t in names:
                tbl = schema.table(t)
                z.writestr(f"{t}.csv", _csv(dataset[t], [c.name for c in tbl.columns]))  # type: ignore[union-attr]
            z.writestr("schema_manifest.json", json.dumps({"database": schema.database, "dbType": schema.db_type, "tables": [t.model_dump(by_alias=True, exclude_none=True) for t in schema.tables if t.name in dataset], "loadOrder": names}, indent=2))
            z.writestr("validation_report.json", report.model_dump_json(by_alias=True, indent=2))
        return buf.getvalue(), f"{STAMP}_data.zip", "application/zip"
    if fmt == "json":
        return json.dumps({t: dataset[t] for t in names}, indent=2, default=str).encode(), f"{STAMP}_data.json", "application/json"
    if fmt == "sql-postgres":
        return _sql(schema, dataset, names, lambda s: f'"{s}"', "PostgreSQL", "BEGIN;").encode(), f"{STAMP}_postgres.sql", "text/plain"
    if fmt == "sql-mysql":
        return _sql(schema, dataset, names, lambda s: f"`{s}`", "MySQL", "START TRANSACTION;").encode(), f"{STAMP}_mysql.sql", "text/plain"
    if fmt == "testcase-xlsx":
        return _xlsx_bytes([("Test cases", _case_rows(cases)), ("Endpoint coverage", _coverage_rows(cases))]), f"{STAMP}_test_cases.xlsx", XLSX
    if fmt == "mapping-report":
        return _xlsx_bytes([("Endpoint coverage", _coverage_rows(cases))]), f"{STAMP}_endpoint_coverage.xlsx", XLSX
    if fmt == "postman":
        return json.dumps(_postman(schema, cases, api_base), indent=2).encode(), f"{STAMP}_postman_collection.json", "application/json"
    if fmt == "pytest":
        return _pytest(cases).encode(), f"test_{STAMP}.py", "text/x-python"
    if fmt == "report-json":
        return report.model_dump_json(by_alias=True, indent=2).encode(), f"{STAMP}_report.json", "application/json"
    if fmt == "report-xlsx":
        r = report
        summary = [
            ("Database type", r.db_type), ("Database", r.database), ("Tables discovered", r.tables_discovered), ("Generation mode", r.mode),
            ("Records requested", r.records_requested), ("Records generated", r.records_generated), ("Data-type validity %", r.data_type_validity),
            ("Relationship integrity %", r.relationship_integrity), ("Uniqueness violations", r.uniqueness_violations), ("Status", r.status), ("Duration (ms)", r.duration_ms),
        ]
        sheets = [("Summary", [{"Metric": k, "Value": v} for k, v in summary]), ("Checks", [c.model_dump() for c in r.checks]), ("Per table", [p.model_dump() for p in r.per_table])]
        if r.scenario_coverage:
            sheets.append(("Scenario coverage", [{"Table": c.scenario.table, "Column": c.scenario.column, "Value": c.scenario.value, "Requested": c.scenario.count, "Generated": c.generated, "Result": c.result, "Note": c.note or ""} for c in r.scenario_coverage]))
        return _xlsx_bytes(sheets), f"{STAMP}_report.xlsx", XLSX
    raise ValueError(f"Unknown export format: {fmt}")
